"""Read a Bibit transaction report (CSV or XLSX) into ledger entries.

The export is Indonesian-formatted and quoted with padding, e.g.

    26 Jul 2023,Majoris Pasar Uang Syariah,Reksadana," 300.000,0 ",...,Pembelian,...

so both the amount parser and the date parser are deliberately forgiving: they
accept ID and EN number formats and ID and EN month abbreviations, and report
what they could not read instead of silently coercing it to zero.
"""

from __future__ import annotations

import csv
import datetime as _dt
import re
from dataclasses import dataclass, field
from pathlib import Path

from .model import IN, OUT, Entry, normalize_month
from .records import RawRow, Report, normalize_kind

DEFAULT_CATEGORY = "Bibit"

#: Bibit transaction verbs -> which way the money moved, from the portfolio's side.
#: "Kupon" (bond coupon) is cash the bond pays out. It is mapped to IN here to
#: match how the original script counted it -- as money staying in the portfolio,
#: reinvested. It is imported as its own row, so flipping it to Out is one edit
#: away if you would rather treat coupons as a withdrawal.
TRANSACTION_DIRECTION = {
    "pembelian": IN,
    "subscription": IN,
    "top up": IN,
    "topup": IN,
    "switch in": IN,
    "kupon": IN,
    "coupon": IN,
    "dividen": IN,
    "dividend": IN,
    "penjualan": OUT,
    "redemption": OUT,
    "penarikan": OUT,
    "withdrawal": OUT,
    "switch out": OUT,
}

_MONTH_NAMES = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "mei": 5, "jun": 6,
    "jul": 7, "aug": 8, "agu": 8, "agt": 8, "sep": 9, "sept": 9, "oct": 10,
    "okt": 10, "nov": 11, "dec": 12, "des": 12,
}
_DMY_RE = re.compile(r"^\s*(\d{1,2})[\s/-]+([A-Za-z]{3,9})[\s/-]+(\d{4})\s*$")
_ISO_RE = re.compile(r"^\s*(\d{4})-(\d{1,2})-(\d{1,2})\b")
_NUMERIC_DMY_RE = re.compile(r"^\s*(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})\s*$")
_NUMERIC_RE = re.compile(r"[^\d,.\-]")

# Header aliases, so a differently-labelled export still lands.
_COLUMNS = {
    "date": ("tanggal transaksi", "tanggal", "date", "transaction date"),
    "product": ("produk", "product", "nama produk"),
    "product_type": ("tipe produk", "jenis produk", "product type"),
    "amount": ("jumlah investasi", "jumlah", "amount", "nominal"),
    "transaction": ("transaksi", "tipe transaksi", "transaction", "type"),
    "currency": ("currency", "mata uang"),
}


class ImportError_(ValueError):
    """Raised when the file cannot be read as a Bibit report at all."""


@dataclass
class ImportResult:
    entries: list = field(default_factory=list)
    months: list = field(default_factory=list)
    rows_read: int = 0
    rows_used: int = 0
    warnings: list = field(default_factory=list)
    sources: list = field(default_factory=list)
    stitch: object = None      # a StitchResult when several files were joined

    @property
    def rows_skipped(self) -> int:
        return self.rows_read - self.rows_used

    def describe(self) -> str:
        if not self.entries:
            return f"Nothing imported from {self.rows_read} rows."
        span = f"{self.months[0]} to {self.months[-1]}" if self.months else "no months"
        if len(self.sources) > 1:
            files = f"{len(self.sources)} files"
        elif self.sources:
            files = self.sources[0]
        else:
            files = "1 file"
        return (
            f"{files}: {self.rows_used} of {self.rows_read} transactions -> "
            f"{len(self.entries)} ledger rows, {span}."
        )


# --------------------------------------------------------------------------- parsing


def parse_amount(raw):
    """Parse an amount written in either Indonesian or English convention.

    '300.000,0' -> 300000.0 ; '1,234.56' -> 1234.56 ; '-' and '' -> None.
    When both separators appear, the rightmost one is the decimal point.
    """
    if raw is None:
        return None
    if isinstance(raw, bool):
        return None
    if isinstance(raw, (int, float)):
        return None if raw != raw else float(raw)

    text = _NUMERIC_RE.sub("", str(raw).strip())
    if not text or text in {"-", ".", ","}:
        return None

    negative = text.startswith("-")
    text = text.lstrip("-")
    last_dot, last_comma = text.rfind("."), text.rfind(",")

    if last_dot >= 0 and last_comma >= 0:
        decimal = "," if last_comma > last_dot else "."
        group = "." if decimal == "," else ","
        text = text.replace(group, "").replace(decimal, ".")
    elif last_comma >= 0:
        parts = text.split(",")
        # A lone comma with a non-3-digit tail is a decimal comma (1234,5).
        text = text.replace(",", ".") if len(parts) == 2 and len(parts[1]) != 3 else text.replace(",", "")
    elif last_dot >= 0:
        parts = text.split(".")
        if not (len(parts) == 2 and len(parts[1]) != 3):
            text = text.replace(".", "")  # thousands separator (300.000)

    try:
        value = float(text)
    except ValueError:
        return None
    return -value if negative else value


def parse_date(raw):
    """Return a datetime.date for a Bibit date cell, or None if it is unreadable.

    Accepts '26 Jul 2023' and '15 Des 2025' (English and Indonesian month names,
    as either appears depending on the app's language) and ISO dates.
    """
    if raw is None:
        return None
    if isinstance(raw, _dt.datetime):
        return raw.date()
    if isinstance(raw, _dt.date):
        return raw
    if hasattr(raw, "year") and hasattr(raw, "month") and hasattr(raw, "day"):
        return _dt.date(int(raw.year), int(raw.month), int(raw.day))

    text = str(raw).strip()
    if not text or text == "-":
        return None

    match = _DMY_RE.match(text)
    if match:
        token = match.group(2).lower()
        month = _MONTH_NAMES.get(token[:4]) or _MONTH_NAMES.get(token[:3])
        if month:
            try:
                return _dt.date(int(match.group(3)), month, int(match.group(1)))
            except ValueError:
                return None

    match = _ISO_RE.match(text)
    if match:
        try:
            return _dt.date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
        except ValueError:
            return None

    # All-numeric dates are read day-first, the Indonesian convention.
    match = _NUMERIC_DMY_RE.match(text)
    if match:
        try:
            return _dt.date(int(match.group(3)), int(match.group(2)), int(match.group(1)))
        except ValueError:
            return None
    return None


def parse_month(raw):
    """Return 'YYYY-MM' for a Bibit date cell, or None if it is unreadable."""
    if raw is None:
        return None
    if hasattr(raw, "year") and hasattr(raw, "month") and not isinstance(raw, str):
        return normalize_month(raw)
    when = parse_date(raw)
    return normalize_month(when) if when else None


# --------------------------------------------------------------------------- reading


def _normalize_header(name) -> str:
    return re.sub(r"\s+", " ", str(name or "").strip().lower())


def _resolve_columns(headers):
    """Map our field names onto whatever this export actually calls its columns."""
    seen = {_normalize_header(h): h for h in headers}
    resolved = {}
    for field_name, aliases in _COLUMNS.items():
        for alias in aliases:
            if alias in seen:
                resolved[field_name] = seen[alias]
                break
    missing = [f for f in ("date", "amount", "transaction") if f not in resolved]
    if missing:
        raise ImportError_(
            "This does not look like a Bibit transaction report - could not find "
            f"the {', '.join(missing)} column(s). Found: {', '.join(map(str, headers))}"
        )
    return resolved


def _read_rows(path: Path):
    """Yield dict rows from a .csv or .xlsx export."""
    if path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}:
        # openpyxl directly rather than via pandas: the sheet is a plain table,
        # and pandas would be 55 MB of dependency for one call.
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ImportError_(
                "Reading .xlsx needs the 'openpyxl' package - run: pip install openpyxl"
            ) from None
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as error:
            raise ImportError_(f"{path.name} could not be opened as a workbook: {error}") from error
        try:
            sheet = workbook.active
            stream = sheet.iter_rows(values_only=True)
            headers = [str(cell) if cell is not None else "" for cell in next(stream, ())]
            rows = [
                dict(zip(headers, values))
                for values in stream
                if any(value is not None and str(value).strip() for value in values)
            ]
        finally:
            workbook.close()
        return headers, rows

    with path.open(newline="", encoding="utf-8-sig") as handle:
        sample = handle.read(8192)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(handle, dialect=dialect)
        headers = list(reader.fieldnames or [])
        return headers, list(reader)


#: Which section of the new PDF layout an old CSV row belongs to, so rows from
#: either format describe themselves the same way once they are stitched together.
_CSV_SECTIONS = {
    "reksadana": ("mutual_fund", "Mutual fund", "Fund"),
    "obligasi fr": ("bond", "FR bonds", "Bond"),
    "obligasi": ("bond", "FR bonds", "Bond"),
    "saham": ("stock", "Stocks", "Stock"),
}
_COUPON_SECTION = ("bond_coupon", "FR bond coupons", "Coupon")
_DIVIDEND_SECTION = ("stock_dividend", "Stock dividends", "Dividend")


def read_tabular_report(path) -> Report:
    """Read a Bibit CSV/XLSX export into a Report of RawRows.

    The spreadsheet export never printed the period it covers, so the range is
    inferred from the rows themselves. That is treated as slightly less
    trustworthy at the edges when stitching -- see Report.interiority.
    """
    path = Path(path)
    if not path.exists():
        raise ImportError_(f"No such file: {path}")

    headers, rows = _read_rows(path)
    columns = _resolve_columns(headers)
    report = Report(path=path, kind=path.suffix.lower().lstrip("."), period_inferred=True)
    unknown_kinds = set()

    for number, row in enumerate(rows, start=2):
        when = parse_date(row.get(columns["date"]))
        amount = parse_amount(row.get(columns["amount"]))
        kind_raw = str(row.get(columns["transaction"]) or "").strip()
        direction = TRANSACTION_DIRECTION.get(kind_raw.lower())

        if when is None:
            report.warnings.append(f"row {number}: unreadable date {row.get(columns['date'])!r}")
            continue
        if amount is None or amount == 0:
            report.warnings.append(f"row {number}: unreadable amount {row.get(columns['amount'])!r}")
            continue
        if direction is None:
            unknown_kinds.add(kind_raw or "(blank)")
            continue

        product_type = str(row.get(columns.get("product_type", "")) or "").strip()
        section = _CSV_SECTIONS.get(product_type.lower(), _CSV_SECTIONS["reksadana"])
        kind = normalize_kind(kind_raw)
        if kind == "Coupon":
            section = _COUPON_SECTION
        elif kind == "Dividend":
            section = _DIVIDEND_SECTION

        amount = abs(amount)
        report.rows.append(RawRow(
            date=when,
            section=section[0], section_label=section[1], short_label=section[2],
            name=str(row.get(columns.get("product", "")) or "").strip(),
            kind=kind,
            inflow=amount if direction == IN else 0.0,
            outflow=amount if direction == OUT else 0.0,
        ))

    if unknown_kinds:
        report.warnings.append(
            "skipped unrecognised transaction type(s): " + ", ".join(sorted(unknown_kinds))
        )
    if not report.rows:
        raise ImportError_(f"{path.name}: no usable transactions found.")

    dates = [row.date for row in report.rows]
    report.period_start, report.period_end = min(dates), max(dates)
    report.rows.sort(key=lambda row: (row.date, row.section, row.name))
    return report


def read_report(path) -> Report:
    """Read any Bibit export -- .pdf, .csv or .xlsx -- into a Report."""
    path = Path(path)
    if path.suffix.lower() == ".pdf":
        # imported late so pdfplumber stays optional unless a PDF is actually used
        from .pdf_report import PdfReportError, parse_report

        try:
            return parse_report(path)
        except PdfReportError as error:
            raise ImportError_(str(error)) from error
    return read_tabular_report(path)


def entries_from_rows(rows, summarize=True, split_by_product_type=False,
                      category=DEFAULT_CATEGORY):
    """Turn RawRows into ledger entries.

    summarize=True (the default) collapses each month to one row per
    (direction, category, section, kind) -- so a month shows up as e.g. a
    Fund Buy row, a Coupon row and a Fund Sell row, all editable. Set it False
    to keep every individual transaction as its own row.
    """
    entries = []
    buckets = {}
    for row in rows:
        name = f"{category} {row.section_label}" if split_by_product_type else category
        month = normalize_month(row.date)
        for direction, amount in ((IN, row.inflow), (OUT, row.outflow)):
            if not amount:
                continue
            if summarize:
                key = (month, direction, name, row.section, row.label)
                slot = buckets.setdefault(key, {"amount": 0.0, "count": 0, "products": set()})
                slot["amount"] += amount
                slot["count"] += 1
                if row.name:
                    slot["products"].add(row.name)
            else:
                entries.append(Entry(
                    month=month, direction=direction, amount=amount, category=name,
                    label=row.label, note=row.name, source="import",
                ))

    for (month, direction, name, _section, label), slot in sorted(buckets.items()):
        products = sorted(slot["products"])
        if not products:
            note = ""
        elif len(products) == 1:
            note = products[0]
        else:
            note = f"{products[0]} +{len(products) - 1} more"
        entries.append(Entry(
            month=month, direction=direction, amount=slot["amount"], category=name,
            label=f"{label} ({slot['count']}x)" if slot["count"] > 1 else label,
            note=note, source="import",
        ))
    return entries


def import_paths(paths, summarize=True, split_by_product_type=False,
                 category=DEFAULT_CATEGORY) -> ImportResult:
    """Read one or more Bibit exports and stitch them into a single set of entries.

    Passing several files is the normal case now that a PDF report covers at most
    a year. They are joined by stitch.stitch(), which takes each day from exactly
    one file, so overlapping reports cannot double-count. Mixing an old CSV export
    with new PDFs works for the same reason.
    """
    from .stitch import StitchError, stitch

    paths = [Path(p) for p in ([paths] if isinstance(paths, (str, Path)) else list(paths))]
    if not paths:
        raise ImportError_("No files to import.")

    reports, failures = [], []
    for path in paths:
        try:
            reports.append(read_report(path))
        except ImportError_ as error:
            failures.append(str(error))
    if not reports:
        raise ImportError_("\n".join(failures) or "Nothing could be read.")

    try:
        joined = stitch(reports)
    except StitchError as error:
        raise ImportError_(str(error)) from error

    result = ImportResult(
        rows_read=joined.rows_in,
        rows_used=len(joined.rows),
        entries=entries_from_rows(joined.rows, summarize, split_by_product_type, category),
        warnings=failures + list(joined.warnings),
        sources=[report.name for report in joined.reports],
        stitch=joined,
    )
    result.months = sorted({entry.month for entry in result.entries})
    return result


def import_file(path, summarize=True, split_by_product_type=False, category=DEFAULT_CATEGORY):
    """Import a single Bibit export. Kept for callers that only ever have one file."""
    return import_paths([path], summarize, split_by_product_type, category)
